"""Explicit TTM and missing-data regression at shared calculation boundaries."""
import json
from types import SimpleNamespace

import pytest
from pydantic import ValidationError
from engine.data_dictionary import (RawFinancials, PreparedTTM, CompanyValuationInput,
    MacroInputs, IndustryData, MethodologyChoices, ValuationAssumptions, QuarterlyFinancials)
from engine.company_metrics import compute_revenue_growth
from engine.ltm_calculator import compute_ltm_financials
from engine.orchestrator import run_full_valuation


def payload():
    raw = RawFinancials(fiscal_year=2026, revenues=100, ebit=20, capex=5, d_a=2,
        net_income=15, bv_equity=40, bv_debt=10, cash_and_marketable_securities=3,
        minority_interests=0, cross_holdings=0, shares_outstanding=10)
    ttm = PreparedTTM(financials=raw, period_start='2025-07-01', period_end='2026-06-30',
        information_as_of='2026-09-06T00:00:00Z', currency='CNY',
        money_unit='million_reporting_currency', shares_unit='million_shares',
        provenance={'source': 'explicit test fixture'})
    return CompanyValuationInput(ticker='TEST', reporting_currency='CNY', prepared_ttm=ttm,
        macro_inputs=MacroInputs(risk_free_rate=.03, equity_risk_premium=.06, tax_rate_marginal=.25),
        industry_data=IndustryData(industry_name='unused', beta_u=0),
        methodology_choices=MethodologyChoices(cost_of_capital_approach='direct', wacc_direct_input=.09),
        valuation_assumptions=ValuationAssumptions(revenue_growth_next_year=.1))


def test_prepared_roundtrip_missing_and_explicit_zero():
    inputs = CompanyValuationInput.model_validate_json(payload().model_dump_json())
    report = run_full_valuation(inputs)
    assert report.ltm_financials.revenues == 100
    assert report.cashflow.fcff is None and report.cashflow.fcfe is None
    assert report.multiples.pe_ratio_intrinsic is None
    assert report.multiples.ev_sales_intrinsic is None
    assert len(report.dcf.fcff_projections) == 10
    assert report.warnings
    inputs.prepared_ttm.financials.change_in_noncash_wc = 0
    inputs.prepared_ttm.financials.net_debt_issued = 0
    report = run_full_valuation(inputs)
    assert report.cashflow.fcff == 12
    assert report.cashflow.fcfe == 12


@pytest.mark.parametrize('changes', [dict(quarters_since_10k=2), dict(reporting_currency='USD')])
def test_prepared_rejects_rotation_and_currency(changes):
    data = payload().model_dump() | changes
    with pytest.raises(ValidationError):
        CompanyValuationInput.model_validate(data)


@pytest.mark.parametrize('changes', [dict(period_start='2026-01-01'),
    dict(information_as_of='2026-06-01T00:00:00Z'), dict(information_as_of='2026-09-06T00:00:00'),
    dict(money_unit='CNY'), dict(provenance={})])
def test_prepared_metadata(changes):
    with pytest.raises(ValidationError):
        PreparedTTM.model_validate(payload().prepared_ttm.model_dump() | changes)


def test_quarterly_missing():
    annual = RawFinancials(fiscal_year=2025, revenues=100, ebit=20, capex=8, bv_debt=10)
    quarter = RawFinancials(fiscal_year=2026, revenues=30, ebit=6)
    qs = [quarter] * 6
    assert compute_ltm_financials(annual, qs, 2).capex is None
    assert compute_ltm_financials(annual, qs, 2).bv_debt is None
    with pytest.raises(ValueError, match='insufficient'):
        compute_ltm_financials(annual, qs[:2], 2)
    complete = [quarter.model_copy(update={'capex': 0}) for _ in range(6)]
    assert compute_ltm_financials(annual, complete, 2).capex == 8


@pytest.mark.parametrize('field', ['bv_debt', 'cash_and_marketable_securities',
    'minority_interests', 'cross_holdings', 'shares_outstanding'])
def test_missing_bridge_does_not_publish_price(field):
    inputs = payload()
    setattr(inputs.prepared_ttm.financials, field, None)
    report = run_full_valuation(inputs)
    assert report.dcf.value_of_operating_assets is not None
    assert report.dcf.value_per_share_pre_options is None
    assert report.final.value_per_share is None


def test_ttm_growth_requires_aligned_comparison_and_does_not_drive_forecast():
    inputs = payload()
    inputs.raw_financials = [inputs.prepared_ttm.financials.model_copy(update={'fiscal_year':2025, 'revenues':80})]
    before = run_full_valuation(inputs)
    assert inputs.company_metrics.revenue_growth is None  # 100/80-1不是TTM同比。
    assert '缺上年同截止月' in inputs.company_metrics.revenue_growth_basis
    evidence = dict(period_end='2025-06-30', currency='CNY', money_unit='million_reporting_currency', revenues='50')
    inputs.prepared_ttm.provenance['prior_ttm_revenue'] = json.dumps(evidence)
    after = run_full_valuation(inputs)
    assert inputs.company_metrics.revenue_growth == 1
    assert after.dcf == before.dcf and after.final == before.final
    inputs.prepared_ttm.financials.revenues = 0
    assert compute_revenue_growth(inputs, inputs.prepared_ttm.financials)[0] == -1
    for change in [dict(revenues='0'), dict(revenues='-1'), dict(revenues='NaN'),
                   dict(currency='USD'), dict(period_end='2025-12-31'), dict(money_unit='CNY')]:
        inputs.prepared_ttm.provenance['prior_ttm_revenue'] = json.dumps(evidence | change)
        assert compute_revenue_growth(inputs, inputs.prepared_ttm.financials)[0] is None


@pytest.mark.parametrize('k', [0,1,2,3,4])
def test_annual_and_quarterly_growth_use_one_year_apart_windows(k):
    inputs = payload()
    inputs.prepared_ttm = None
    inputs.raw_financials = [RawFinancials(fiscal_year=2025, revenues=100, ebit=20),
                            RawFinancials(fiscal_year=2024, revenues=50, ebit=10)]
    inputs.quarters_since_10k = k
    inputs.quarterly_financials = [QuarterlyFinancials(fiscal_year=2026, revenues=v, ebit=v*.2)
        for v in [40,35,30,25,20,15,10,5]]
    current = compute_ltm_financials(inputs.raw_financials[0], inputs.quarterly_financials, k)
    growth, _ = compute_revenue_growth(inputs, current)
    assert growth == current.revenues / 50 - 1
    if k:
        for invalid in [None, float('inf'), float('nan'), 1e308]:
            inputs.quarterly_financials[7].revenues = invalid
            if invalid == 1e308:
                inputs.quarterly_financials[6].revenues = invalid
            assert compute_revenue_growth(inputs, current)[0] is None
    else:
        inputs.raw_financials[1].fiscal_year = 2023
    assert compute_revenue_growth(inputs, current)[0] is None


def test_prepared_growth_handles_leap_day_without_blocking_valuation():
    from datetime import date
    inputs = payload()
    inputs.prepared_ttm.period_start = date(2023,3,1)
    inputs.prepared_ttm.period_end = date(2024,2,29)
    inputs.prepared_ttm.provenance['prior_ttm_revenue'] = json.dumps(dict(
        period_end='2023-02-28', currency='CNY', money_unit='million_reporting_currency', revenues='50'))
    report = run_full_valuation(inputs)
    assert inputs.company_metrics.revenue_growth == 1 and report.dcf is not None


def test_excel_ltm_columns_use_engine_flows_and_latest_balances():
    from openpyxl import Workbook
    from api.export import _build_input_sheet
    from api.export_workbook import _write_input_sheet, _write_trailing_twelve_month, _write_diagnostics
    inputs = payload()
    inputs.raw_financials = [inputs.prepared_ttm.financials.model_copy(update={'fiscal_year':2025, 'revenues':80, 'bv_debt':70, 'cash_and_marketable_securities':50})]
    report = run_full_valuation(inputs)
    wb = Workbook()
    _write_input_sheet(wb, inputs, report)
    assert wb['Input Sheet']['B12'].value == 100
    assert wb['Input Sheet']['C12'].value == 80
    assert wb['Input Sheet']['B20'].value == 10
    assert wb['Input Sheet']['B23'].value == 3
    _write_trailing_twelve_month(wb, inputs, report)
    ws = wb['Trailing 12 Month']
    assert [ws.cell(3, c).value for c in (2,3,4)] == [100,80,100]
    _write_diagnostics(wb, inputs, report)
    ws = wb['Diagnostics']
    assert ws['B4'].value is None
    assert next(row[1].value for row in ws if row[0].value == 'Base year revenues') == 100
    simple = _build_input_sheet(SimpleNamespace(inputs=inputs, report=report)).active
    assert next(row[1].value for row in simple if row[0].value == 'Revenues') == 100
    inputs.prepared_ttm = None
    inputs.quarters_since_10k = 2
    inputs.quarterly_financials = [QuarterlyFinancials(fiscal_year=2026, revenues=20, ebit=4, bv_debt=7, net_income=0) for _ in range(8)]
    inputs.raw_financials[0].net_income = 0
    report = run_full_valuation(inputs)
    wb = Workbook()
    _write_input_sheet(wb, inputs, report)
    assert wb['Input Sheet']['B20'].value == 7
    assert wb['Input Sheet']['B15'].value == 0


@pytest.mark.parametrize('negative', [False, True])
def test_full_excel_is_shared_engine_snapshot_not_fixed_five_year_calculator(negative):
    import io
    from openpyxl import load_workbook
    from api.export_workbook import generate_full_workbook
    inputs = payload()
    inputs.valuation_assumptions.high_growth_years = 1
    inputs.valuation_assumptions.projection_years = 7
    inputs.valuation_assumptions.revenue_growth_years_2_5 = .1
    inputs.valuation_assumptions.target_operating_margin = -.05 if negative else .3
    inputs.valuation_assumptions.override_nol = True
    inputs.valuation_assumptions.nol_amount = 200
    inputs.valuation_assumptions.stable_growth_rate = 0  # 显式0不回退Rf。
    inputs.valuation_assumptions.sales_to_capital_high = 4
    inputs.valuation_assumptions.sales_to_capital_stable = 5
    inputs.valuation_assumptions.failure_probability = .2
    report = run_full_valuation(inputs)
    wb = load_workbook(io.BytesIO(generate_full_workbook(inputs, report)))
    assert len(wb.sheetnames) == 13
    ws = wb['Valuation Output']
    assert [ws.cell(4, c).value for c in range(3, 10)] == pytest.approx(report.dcf.revenue_projections, rel=1e-15)
    assert [ws.cell(8, c).value for c in range(3, 10)] == pytest.approx(report.dcf.nopat_projections, rel=1e-15)
    assert [ws.cell(9, c).value for c in range(3, 10)] == pytest.approx(report.dcf.reinvestment_projections, rel=1e-15)
    assert [ws.cell(10, c).value for c in range(3, 10)] == pytest.approx(report.dcf.fcff_projections, rel=1e-15)
    assert [ws.cell(12, c).value for c in range(3, 10)] == pytest.approx(report.dcf.discount_factors, rel=1e-15)
    assert ws['J2'].value == 'Terminal year' and ws['J10'].value == pytest.approx(report.dcf.terminal_fcff)
    assert ws['J13'].value == pytest.approx(report.dcf.terminal_value_firm)
    assert ws['B18'].value == pytest.approx(report.dcf.value_of_operating_assets)
    assert ws['B27'].value == pytest.approx(report.final.value_per_share)
    assert wb['Cost of Capital']['B28'].value == report.cost_of_capital.wacc  # direct分支不重算Beta。
    assert wb['Input Sheet']['B36'].value == report.dcf.forecast_parameters['operating_margin_next_year'].value
    assert wb['Input Sheet']['A37'].value == 'Revenue Growth - Transition Anchor'
    assert '快照' in ws['A31'].value
    assert all(cell.data_type != 'f' for name in ['Valuation Output', 'Cost of Capital'] for row in wb[name] for cell in row)
