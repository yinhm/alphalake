"""年度总投入沿实际投资年：独立现金流、资本桥及API拒绝保护。"""
from copy import deepcopy
import json
import math

import pytest
from pydantic import ValidationError

from engine.data_dictionary import CompanyValuationInput, ValuationAssumptions
from engine.module_4_dcf import InvalidReinvestment, _reinvestment_path
from engine.orchestrator import run_full_valuation
from tests.test_native_assumption_selection import sample
from tools.review_native_policy import review_report


def inputs_for(growth=0, research=True, lag=1):
    inputs = CompanyValuationInput.model_validate(sample()['inputs'])
    inputs.adjustment_inputs.has_r_and_d = research
    inputs.adjustment_inputs.r_and_d_expense_current = 100
    inputs.adjustment_inputs.r_and_d_expense_past = [40]*5
    for i, row in enumerate(inputs.raw_financials):
        row.r_and_d_expense = 100 if i == 0 else 40
    inputs.valuation_assumptions = ValuationAssumptions(
        projection_years=3, high_growth_years=3,
        annual_forecast=[dict(growth=growth, margin=.2, tax=.25)]*3,
        stable_growth_rate=.02, cost_of_capital_stable_override=.08, roic_stable_override=.08,
        override_reinvestment_lag=True, reinvestment_lag_years=lag,
        research_tax=(dict(policy_id='synthetic-test', expenses=[100]*3, deductions=[100]*3,
            terminal_expense=100, terminal_deduction=100, tax_rate_basis='taxable_profit_rate')
            if research else None))
    return inputs


def body_for(inputs, report):
    return dict(inputs=inputs.model_dump(mode='json'),
        **{k:getattr(report,k).model_dump(mode='json') for k in
           ('ltm_financials','adjusted','cost_of_capital','cashflow','dcf','final')},
        convertible_claim=report.convertible_claim, equity_bridge=report.equity_bridge)


@pytest.mark.parametrize('growth', [0, .1, -.1])
@pytest.mark.parametrize('research', [False, True])
@pytest.mark.parametrize('lag', [0, 1, 2, 3])
def test_total_budget_changes_only_investment_cashflow_and_capital(growth, research, lag):
    inputs = inputs_for(growth, research, lag)
    baseline = run_full_valuation(inputs)
    changed = inputs.model_copy(deep=True)
    budget = [48, 0, -5]
    changed.valuation_assumptions.annual_net_reinvestment = budget
    report = run_full_valuation(changed)
    dcf = report.dcf
    assert dcf.reinvestment_projections == budget
    for field in ('ltm_financials','adjusted','cost_of_capital','cashflow'):
        assert getattr(report, field) == getattr(baseline, field)
    for field, value in baseline.dcf.model_dump().items():
        if field not in ('forecast_parameters','reinvestment_projections','fcff_projections','pv_fcff',
                'pv_cash_flows_sum','value_of_operating_assets','value_of_equity','value_per_share_pre_options',
                'implied_roic_projections','implied_roic_terminal'):
            assert getattr(dcf, field) == getattr(baseline.dcf, field), field
    assert dcf.fcff_projections == pytest.approx([p-b for p,b in zip(dcf.nopat_projections, budget)])
    pv_change = -sum((b-old)*df for b,old,df in zip(
        budget, baseline.dcf.reinvestment_projections, dcf.discount_factors))
    assert dcf.value_of_operating_assets-baseline.dcf.value_of_operating_assets == pytest.approx(pv_change)
    assert report.final.value_per_share-baseline.final.value_per_share == pytest.approx(pv_change/10)
    capital = (report.ltm_financials.consolidated_book_equity+report.ltm_financials.bv_debt
        -report.ltm_financials.cash_and_marketable_securities+report.adjusted.value_of_research_asset
        +report.adjusted.pv_of_operating_leases)
    for nopat, investment, roic in zip(dcf.nopat_projections, budget, dcf.implied_roic_projections):
        assert roic == pytest.approx(nopat/capital)
        capital += investment
    assert dcf.implied_roic_terminal == pytest.approx(dcf.terminal_nopat/capital)
    # Unused scalar ratios and lag cannot silently move or change the budget.
    changed.valuation_assumptions.sales_to_capital_high = 0
    changed.valuation_assumptions.sales_to_capital_stable = -2
    changed.valuation_assumptions.reinvestment_lag_years = 3-lag
    assert run_full_valuation(changed).dcf == dcf
    body = body_for(changed, report)
    reviewed = review_report(body)
    assert reviewed['capital']['annual_net_reinvestment'] == budget
    assert not reviewed['capital']['reinvestment_lag_applied']
    assert not reviewed['capital']['budget_evidence_verified']
    for row in reviewed['forecast']:
        assert row['capital_funding']['revenue_year'] is None
        assert row['capital_funding']['sales_to_capital'] is None
        assert row['incremental_return_bridge']['status'] == 'not_inferred_from_explicit_budget'
    transition = reviewed['economic_checks']['terminal_transition']
    assert transition['continued_capital_ratio_reinvestment_million_cny'] is None
    assert transition['continued_explicit_budget_reinvestment_million_cny'] == pytest.approx(-5*1.02)
    assert transition['requires_capital_transition_basis']
    if research:
        assert dcf.research_amortization_projections == [52, 64, 76]
        assert reviewed['forecast'][0]['implied_nonresearch_net_reinvestment_million_cny'] == 0
        terminal = reviewed['terminal']['research_capital']
        assert terminal['net_investment_million_cny']+terminal['nonresearch_net_investment_million_cny'] == pytest.approx(dcf.terminal_reinvestment)
    tampered = deepcopy(body); tampered['dcf']['reinvestment_projections'][0] += 1
    with pytest.raises(ValueError, match='saved report does not replay'):
        review_report(tampered)


def test_budget_validation_and_unprovided_path_are_not_changed():
    for values in ([], [0], [0]*4, [0, math.nan, 0], [0, math.inf, 0], [0, -math.inf, 0]):
        with pytest.raises(ValidationError):
            ValuationAssumptions(projection_years=3, annual_net_reinvestment=values)
        with pytest.raises(InvalidReinvestment):
            _reinvestment_path([], 0, 0, 3, 1, 3, annual_net_reinvestment=values)
    with pytest.raises(ValidationError, match='mutually exclusive'):
        ValuationAssumptions(projection_years=3, annual_net_reinvestment=[0]*3, annual_sales_to_capital=[2]*3)
    with pytest.raises(InvalidReinvestment, match='mutually exclusive'):
        _reinvestment_path([], 2, 2, 0, 1, 3, [2]*3, [0]*3)
    assert _reinvestment_path([100, 90, 90, 100], 2, 4, 0, 1, 3) == [-5, 0, 2.5]
    assert _reinvestment_path([100, 90, 90, 100], 2, 4, 0, 1, 3, [2, 4, 4]) == [-5, 0, 2.5]


def test_explicit_last_budget_can_align_with_original_terminal_without_changing_roic():
    from tools.cyclical_valuation import economic_review
    inputs = inputs_for()
    before = run_full_valuation(inputs)
    inputs.valuation_assumptions.annual_net_reinvestment = [48, 0, before.dcf.terminal_reinvestment/1.02]
    report = run_full_valuation(inputs)
    assert report.dcf.terminal_roic == before.dcf.terminal_roic == .08
    assert report.dcf.terminal_value_firm == before.dcf.terminal_value_firm
    body = body_for(inputs, report)
    assert not review_report(body)['economic_checks']['terminal_transition']['requires_capital_transition_basis']
    audited = economic_review(body)
    assert audited['terminal_reinvestment_continuous']
    assert not audited['economic_forecast_verified']


def test_api_budget_rejection_preserves_session_and_can_restore_ratio_path(monkeypatch):
    from fastapi.testclient import TestClient
    from api import routes, session_store
    from api.main import app
    from data_sources.damodaran_store import DamodaranStore
    monkeypatch.setattr(routes, '_get_damodaran_store', lambda: DamodaranStore())
    monkeypatch.setattr(routes, '_build_industry_lookup', lambda _: None)
    monkeypatch.setattr(session_store, '_sessions', {})
    with TestClient(app) as client:
        created = client.post('/api/valuation', json={'inputs':inputs_for().model_dump(mode='json')})
        assert created.status_code == 200, created.text
        baseline = created.json(); url = '/api/valuation/'+baseline['id']
        path = 'valuation_assumptions.annual_net_reinvestment'
        response = client.patch(url, json={'overrides':{path:[48, 0, -5]}})
        assert response.status_code == 200, response.text
        saved = response.json()
        assert saved['dcf']['reinvestment_projections'] == [48, 0, -5]
        assert saved['dcf']['forecast_parameters']['sales_to_capital_high']['source'] == 'annual_net_reinvestment'
        assert not saved['dcf']['forecast_parameters']['sales_to_capital_high']['editable']
        assert 'User override' in saved['source_metadata'][path]
        review_report(saved)
        for overrides in ({path:[0]}, {path:[0, math.inf, 0]}, {path:[0, math.nan, 0]},
                {'valuation_assumptions.projection_years':4},
                {'valuation_assumptions.annual_sales_to_capital':[2]*3}):
            rejected = client.patch(url, content=json.dumps({'overrides':overrides}),
                headers={'Content-Type':'application/json'})
            assert rejected.status_code == 422, rejected.text
            assert client.get(url).json() == saved
        sensitivity = client.post(url+'/sensitivity')
        assert sensitivity.status_code == 200, sensitivity.text
        assert 'sales_to_capital_high' not in {r['driver'] for r in sensitivity.json()['bars']}
        assert client.get(url).json() == saved
        restored = client.patch(url, json={'overrides':{path:None}})
        assert restored.status_code == 200, restored.text
        assert restored.json()['dcf'] == baseline['dcf']
        assert restored.json()['final'] == baseline['final']
        # Explicit paths switch atomically; supplying both always rejects.
        assert client.patch(url, json={'overrides':{
            'valuation_assumptions.annual_sales_to_capital':[2]*3}}).status_code == 200
        switched = client.patch(url, json={'overrides':{
            'valuation_assumptions.annual_sales_to_capital':None, path:[0]*3}})
        assert switched.status_code == 200, switched.text
        assert switched.json()['dcf']['reinvestment_projections'] == [0]*3


def test_existing_candidate_generators_do_not_replace_explicit_budget():
    from tools.cyclical_valuation import cycle_conditions
    from tools.compare_native_capital import compare
    from tools.select_native_assumptions import select
    inputs = inputs_for()
    inputs.valuation_assumptions.annual_net_reinvestment = [48, 0, -5]
    body = body_for(inputs, run_full_valuation(inputs)); saved = deepcopy(body)
    with pytest.raises(ValueError, match='explicit_future_path_preserved'):
        cycle_conditions(body)
    assert compare(body)['status'] == 'outside_policy_scope'
    from tests.test_native_assumption_selection import POLICY
    with pytest.raises(ValueError, match='explicit net reinvestment'):
        select(body, body['inputs'], [], '2026-10-05T00:00:00+00:00', POLICY)
    assert body == saved


def test_excel_exports_actual_budget_without_active_scalar_ratios():
    import io
    from openpyxl import load_workbook
    from api.export_workbook import generate_full_workbook
    inputs = inputs_for()
    inputs.valuation_assumptions.annual_net_reinvestment = [48, 0, -5]
    inputs.valuation_assumptions.sales_to_capital_high = 999
    report = run_full_valuation(inputs)
    workbook = load_workbook(io.BytesIO(generate_full_workbook(inputs, report)))
    sheet = workbook['Valuation Output']
    assert len(workbook.sheetnames) == 13
    assert [sheet.cell(9,c).value for c in range(3,6)] == [48, 0, -5]
    assert [sheet.cell(10,c).value for c in range(3,6)] == pytest.approx(report.dcf.fcff_projections)
    assert sheet['F9'].value == pytest.approx(report.dcf.terminal_reinvestment)
    assert workbook['Input Sheet']['B40'].value is None
    for row in workbook['Stories to Numbers']:
        if row[2].value == 'sales_to_capital_high':
            assert row[3].value is None
            assert row[4].value == 'Not consumed: explicit total budget'
